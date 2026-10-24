"""Credential-free regression for network, batch, documents and AI boundaries."""
import base64
import io
import json
import threading
import time
from types import SimpleNamespace

import pytest
import requests
from PIL import Image

import api
import config
import tasks
import question_documents as qd
from desktop.bridge import Bridge, prepare_messages
from desktop.storage import Store, DEFAULTS
from desktop import models


def png():
    buffer = io.BytesIO()
    Image.new('RGB', (20, 10), 'white').save(buffer, 'PNG')
    return buffer.getvalue()


def data_url():
    return 'data:image/png;base64,' + base64.b64encode(png()).decode()


class Response:
    def __init__(self, status=200, payload=None, data=None, headers=None):
        self.status_code = status
        self.payload = payload if payload is not None else {'code': 0, 'data': []}
        self.data = data if data is not None else png()
        self.headers = headers or {'Content-Type': 'image/png'}
    def __enter__(self): return self
    def __exit__(self, *args): self.close()
    def close(self): pass
    def raise_for_status(self):
        if self.status_code >= 400: raise requests.HTTPError('secret-url-token')
    def json(self): return self.payload
    def iter_content(self, size):
        yield self.data


class Session:
    def __init__(self, *responses):
        self.responses = list(responses)
        self.headers = {}
        self.cookies = requests.cookies.RequestsCookieJar()
        self.calls = []
    def __enter__(self): return self
    def __exit__(self, *args): pass
    def close(self): pass
    def mount(self, *args): pass
    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        value = self.responses.pop(0)
        if isinstance(value, Exception): raise value
        return value
    get = post


def client(*responses, url='https://lgxt.wutp.com.cn/api', **kwargs):
    c = api.APIClient(url, **kwargs)
    c.session = Session(*responses)
    return c


def test_default_http_fails_before_network_and_explicit_opt_in():
    c = client(Response(payload={'code':0,'data':'token'}), url='http://lgxt.wutp.com.cn/api')
    assert not c.login('private-user', 'secret-password')[0]
    assert not c.session.calls
    c.allow_insecure = True
    assert c.login('private-user', 'secret-password')[0]
    assert 'Authorization' not in c.session.headers
    c.logout()
    assert 'Authorization' not in c.headers


@pytest.mark.parametrize('method', ['get_my_courses', 'get_user_info'])
def test_reads_retry_finitely_and_honor_retry_after(monkeypatch, method):
    sleeps = []
    monkeypatch.setattr(api.time, 'sleep', sleeps.append)
    c = client(Response(429, headers={'Retry-After':'2'}), Response(503), Response())
    assert getattr(c, method)()[0]
    assert len(c.session.calls) == 3 and sleeps[0] >= 2
    assert all(call[1]['allow_redirects'] is False for call in c.session.calls)
    assert all(call[1]['timeout'] == api.TIMEOUT for call in c.session.calls)
    c = client(Response(429, headers={'Retry-After':'100'}))
    assert not getattr(c, method)()[0] and len(c.session.calls) == 1


def test_read_timeouts_exhaust_budget_and_redact(monkeypatch):
    monkeypatch.setattr(api.time,'sleep',lambda _:None)
    c = client(*[requests.Timeout('token-password-url')] * 3)
    ok, error = c.get_questions(3)
    assert not ok and len(c.session.calls) == 3 and 'token-password' not in error


@pytest.mark.parametrize('response', [requests.Timeout('secret'), Response(503), Response(302), Response(payload=[])])
def test_writes_never_replay_ambiguous_results(response):
    c = client(response)
    ok, error = c.submit_answer(12,100)
    assert not ok and error.startswith('结果待核实') and len(c.session.calls) == 1


def test_definite_rejection_is_not_unknown():
    c = client(Response(429))
    ok, error = c.submit_answer(12,100)
    assert not ok and '429' in error and not error.startswith('结果待核实')


def test_read_cache_suppresses_concurrent_duplicate_lists():
    c = client(Response())
    threads = [threading.Thread(target=c.get_my_courses) for _ in range(10)]
    for t in threads: t.start()
    for t in threads: t.join()
    assert len(c.session.calls) == 1


def test_no_adapter_retries_or_ambient_auth():
    c = api.APIClient('https://lgxt.wutp.com.cn/api')
    assert not c.session.trust_env
    assert c.session.get_adapter(c.base_url).max_retries.total == 0


@pytest.fixture
def public_dns(monkeypatch):
    monkeypatch.setattr(api.socket, 'getaddrinfo', lambda *a,**kw: [(2,1,6,'',('8.8.8.8',443))])


@pytest.mark.parametrize('url', ['file:///etc/passwd','http://lgxt.wutp.com.cn/x','https://evil.example/x',
    'https://lgxt.wutp.com.cn.evil.example/x','https://u:p@lgxt.wutp.com.cn/x',
    'https://lgxt.wutp.com.cn:8443/x','https://127.0.0.1/x','https://[::1]/x',
    'https://169.254.169.254/latest/meta-data/','data:image/png;base64,eA=='])
def test_image_origin_policy(url, public_dns):
    with pytest.raises(ValueError): api.APIClient('https://lgxt.wutp.com.cn/api').fetch_image(url)


@pytest.mark.parametrize('address', ['127.0.0.1','10.0.0.1','192.168.1.2','169.254.169.254','::1','fc00::1'])
def test_image_rejects_private_dns(monkeypatch, address):
    monkeypatch.setattr(api.socket,'getaddrinfo',lambda *a,**kw:[(2,1,6,'',(address,443))])
    with pytest.raises(ValueError): api.APIClient('https://lgxt.wutp.com.cn/api').fetch_image('https://lgxt.wutp.com.cn/img')


def test_images_are_isolated_streamed_and_decoded(monkeypatch, public_dns):
    c = api.APIClient('https://lgxt.wutp.com.cn/api')
    c.headers['Authorization'] = 'PRIVATE'
    images = Session(Response())
    monkeypatch.setattr(api.requests, 'Session', lambda: images)
    assert c.fetch_image('https://lgxt.wutp.com.cn/img') == png()
    assert c.fetch_image('https://lgxt.wutp.com.cn/img') == png()
    assert len(images.calls) == 1
    assert not images.trust_env
    assert 'headers' not in images.calls[0][1] and images.calls[0][1]['stream']
    assert images.calls[0][1]['allow_redirects'] is False


@pytest.mark.parametrize('response', [Response(302), Response(data=b'<svg onload="attack()"/>'),
    Response(headers={'Content-Type':'text/html'}), Response(headers={'Content-Type':'image/jpeg'}), Response(headers={'Content-Type':'image/png','Content-Length':str(api.MAX_IMAGE_BYTES+1)}),
    Response(data=b'x'*(api.MAX_IMAGE_BYTES+1))])
def test_image_rejects_redirect_content_and_sizes(monkeypatch, public_dns, response):
    c=api.APIClient('https://lgxt.wutp.com.cn/api')
    monkeypatch.setattr(api.requests,'Session',lambda:Session(response))
    with pytest.raises(ValueError): c.fetch_image('https://lgxt.wutp.com.cn/img')


def doc():
    return {'schemaVersion':1,'courseId':'1','workId':'2','questionId':'3','title':'Printed formula fixture',
            'courseName':'Fixture course','workName':'Fixture assignment','originalImage':data_url(),
            'sourceText':'Original printed source','recognizedText':'','latex':r'\frac{1}{2}',
            'correctedMarkdown':r'Corrected $\frac{1}{2}$','solutionMarkdown':'## Derivation\n\n**Energy**\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n```python\nx=1\n```','sources':[]}


@pytest.mark.parametrize('formula',[r'\frac{a_1}{b^2}',r'\int_0^1 x^2\,dx',r'\begin{pmatrix}1&0\\0&1\end{pmatrix}',r'\alpha+\beta=\gamma'])
def test_recognition_and_manual_correction_roundtrip(tmp_path, formula):
    parsed=qd.parse_recognition(json.dumps({'text':'Printed text','latex':formula}))
    d=doc()|{'latex':parsed['latex'],'correctedMarkdown':f"User correction $${formula}$$"}
    store=Store(tmp_path)
    saved=store.save_question_document(d)
    assert store.question_document(1,2,3)['correctedMarkdown']==d['correctedMarkdown']
    store.save_question_document(saved|{'solutionMarkdown':'Changed'})
    assert list(tmp_path.glob('question-*.previous.json'))
    assert formula in qd.context_messages(saved,'Solve')[0]['attachments'][0]['text']


@pytest.mark.parametrize('text',['not JSON','[]','{"latex":1,"text":"x"}','{"latex":"x"}', 'x'*64001], ids=['not-json','array','wrong-type','missing-field','too-long'])
def test_malformed_recognition(text):
    with pytest.raises(ValueError): qd.parse_recognition(text)


def test_unsupported_schema_is_not_overwritten(tmp_path):
    store=Store(tmp_path);store.save_question_document(doc())
    with pytest.raises(ValueError): store.save_question_document(doc()|{'schemaVersion':99})
    assert store.question_document(1,2,3)['schemaVersion']==1


def test_provider_vision_capability_and_serialization():
    d=doc(); messages=qd.context_messages(d,'Solve',d['originalImage'])
    for provider in ['OpenAI','Claude','Gemini','Local']:
        s=DEFAULTS|{'provider':provider,'model':'custom-model','visionEnabled':True}
        prepared=prepare_messages(s,messages)
        models.require_vision(s,prepared)
        assert prepared[0]['content'][1]['type']==('image' if provider=='Claude' else 'image_url')
        with pytest.raises(ValueError): models.require_vision(s|{'visionEnabled':False},prepared)
    models.require_vision(DEFAULTS,[{'role':'user','content':'Text only'}])
    with pytest.raises(ValueError): prepare_messages(DEFAULTS,[{'role':'user','attachments':[{'data':'data:image/png;base64,eA=='}]}])


def test_batch_cancellation_duplicates_partial_and_audit(tmp_path):
    cancel=threading.Event();calls=[];active=0
    def submit(wid, grade):
        calls.append(wid)
        if wid==2:
            cancel.set();return False,'结果待核实：timeout'
        return True,'accepted'
    client=SimpleNamespace(submit_answer=submit)
    ui=SimpleNamespace(status=lambda _:None,progress=lambda _:None,maximum=lambda _:None)
    store=Store(tmp_path)
    result=tasks.execute_batch(ui,client,[{'workId':n} for n in [1,1,2,3]],cancel,store.audit_batch)
    assert calls==[1,2] and result['cancelled']
    assert [r['status'] for r in result['results']]==['succeeded','verification_required']
    audit=json.loads((tmp_path/'batch-audit.json').read_text())
    assert len(audit)==2 and 'workId' not in audit[0] and 'message' not in audit[0]


def test_batch_confirmation_snapshot_single_use(tmp_path):
    c=api.APIClient('https://lgxt.wutp.com.cn/api');calls=[]
    c.get_course_works=lambda _: (True,[{'workId':1,'workName':'Fixture'}])
    c.submit_answer=lambda *args:(calls.append(args) or True,'confirmed')
    b=Bridge(tmp_path/'data',client=c,legacy_settings=config.Settings(str(tmp_path/'config.ini')));b._username='fixture'
    plan=b.preview_batch('course',{'courseId':1})
    assert plan['count']==1 and not calls
    jid=b.confirm_batch(plan['token'])
    for _ in range(100):
        if b.poll_job(jid)['done']:break
        time.sleep(.01)
    assert len(calls)==1
    with pytest.raises(ValueError): b.confirm_batch(plan['token'])
    with pytest.raises(ValueError): b.start_task('submit_all')


def test_exports_preserve_source_image_and_equations(tmp_path):
    import exporter
    from docx import Document
    for format in ['md','docx','pdf']:
        d=doc();d['solutionMarkdown']+='\n\n![embedded fixture]('+data_url()+')'
        result=exporter.export_question_document(d,str(tmp_path),format)
        from pathlib import Path
        path=Path(result['path']);assert path.is_file() and path.stat().st_size>100
        assert result['equations']=='latex-source'
        if format=='docx':
            word=Document(path);assert word.tables and len(word.inline_shapes)==2
            assert any('\\frac' in p.text for p in word.paragraphs)
    assert list((tmp_path/'question-documents').glob('*-original.png'))


def test_editor_images_fit_document_limit_without_flooding_ai_context():
    buffer=io.BytesIO()
    Image.effect_noise((512,512),100).save(buffer,'PNG')
    image='data:image/png;base64,'+base64.b64encode(buffer.getvalue()).decode()
    solution='![local diagram]('+image+')'
    assert len(solution)>200000
    d=qd.validate_document(doc()|{'solutionMarkdown':solution})
    assert d['solutionMarkdown']==solution
    context=qd.context_messages(d,'Solve')[0]['attachments'][0]['text']
    assert len(context)<2000 and image not in context
