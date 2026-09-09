# -*- coding: utf-8 -*-
"""导出：题目图片缓存 + Word / PDF 输出。任务线程中调用，不触碰 UI。"""
import collections
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor
from reportlab.lib import colors
from reportlab.lib.units import inch

IMAGE_DIR = '题目图片'
FOLDER_ROOT = '作业'

# 导出开关快照（主线程读取 Tk 变量后传入 worker，避免跨线程访问）
ExportOptions = collections.namedtuple(
    'ExportOptions', 'export_path export_word export_word_answers export_pdf export_pdf_answers')


def clean_name(name):
    return ''.join(c for c in name if c not in r'<>:"/\|?*')


def assignment_folder(export_path, course_name, course_id, work_name, work_id):
    return os.path.join(export_path, FOLDER_ROOT,
                        f'{clean_name(course_name)} (ID_{course_id})',
                        f'{clean_name(work_name)} (ID_{work_id})')


def save_question(client, question, assignment_folder_path):
    """把题目图片下载到 题目图片/<id>.png；已存在且非空则跳过（缓存）。"""
    question_id = question.get('id', 'N/A')
    imgurl = question.get('imgurl', 'N/A')
    images_folder = os.path.join(assignment_folder_path, IMAGE_DIR)
    os.makedirs(images_folder, exist_ok=True)
    if not imgurl or imgurl == 'N/A':
        return
    image_path = os.path.join(images_folder, f'{question_id}.png')
    if os.path.exists(image_path) and os.path.getsize(image_path) > 0:
        return  # 已缓存，避免重复下载
    try:
        with open(image_path, 'wb') as f:
            f.write(client.fetch_image(imgurl))
    except Exception as e:
        print(f'无法下载题目 {question_id} 的图片：{e}')


def save_collected(client, collected_questions, work_name, course_name, work_id, course_id,
                   chapter_name, opts):
    """把收集结果写入磁盘：图片 + 可选 Word/PDF。与旧版目录结构一致。"""
    folder = assignment_folder(opts.export_path, course_name, course_id, work_name, work_id)
    os.makedirs(folder, exist_ok=True)
    for question in collected_questions.values():
        save_question(client, question, folder)
    if opts.export_word:
        save_questions_to_word(collected_questions, folder, chapter_name,
                               export_answers=opts.export_word_answers)
    if opts.export_pdf:
        save_questions_to_pdf(collected_questions, folder, chapter_name,
                              export_answers=opts.export_pdf_answers)


def save_questions_to_word(collected_questions, assignment_folder_path, work_name, export_answers=True):
    document = Document()
    style = document.styles['Normal']
    style.font.name = '宋体'
    style.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    document.add_heading(work_name, 0)

    images_folder = os.path.join(assignment_folder_path, IMAGE_DIR)
    for idx, question_id in enumerate(sorted(collected_questions, key=int)):
        question = collected_questions[question_id]
        document.add_heading(f'题目 {idx + 1}: {question.get("name", "N/A")}', level=2)
        image_path = os.path.join(images_folder, f'{question_id}.png')
        if os.path.exists(image_path):
            document.add_picture(image_path, width=Inches(5))
        else:
            document.add_paragraph('（无图片）')
        if export_answers:
            run = document.add_paragraph('答案：').add_run(question.get('answer', 'N/A'))
            run.font.color.rgb = RGBColor(255, 0, 0)

    document.save(os.path.join(assignment_folder_path, f'{work_name}.docx'))


def save_questions_to_pdf(collected_questions, assignment_folder_path, work_name, export_answers=True):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (BaseDocTemplate, Frame, Image as RLImage,
                                    PageTemplate, Paragraph, Spacer)

    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    styles = getSampleStyleSheet()
    for name, parent, size, leading, color in [
        ('ChineseTitle', 'Title', 20, 24, '#333333'),
        ('ChineseHeading1', 'Heading1', 16, 20, '#555555'),
        ('Chinese', 'Normal', 12, 18, '#000000'),
    ]:
        styles.add(ParagraphStyle(name=name, parent=styles[parent], fontName='STSong-Light',
                                  fontSize=size, leading=leading,
                                  alignment=1 if name == 'ChineseTitle' else 0,
                                  textColor=colors.HexColor(color)))
    normal_style = styles['Chinese']

    def add_page_number(canvas, _doc):
        canvas.setFont('STSong-Light', 9)
        canvas.drawRightString(A4[0] - 50, 15, f'第 {canvas.getPageNumber()} 页')

    doc = BaseDocTemplate(os.path.join(assignment_folder_path, f'{work_name}.pdf'), pagesize=A4,
                          rightMargin=40, leftMargin=40, topMargin=60, bottomMargin=60)
    frame = Frame(40, 60, doc.width, doc.height, id='normal')
    doc.addPageTemplates([PageTemplate(id='test', frames=frame, onPage=add_page_number)])

    elements = [Paragraph(work_name, styles['ChineseTitle']), Spacer(1, 0.3 * inch)]
    images_folder = os.path.join(assignment_folder_path, IMAGE_DIR)
    for idx, question_id in enumerate(sorted(collected_questions, key=int)):
        question = collected_questions[question_id]
        elements.append(Paragraph(f'题目 {idx + 1}: {question.get("name", "N/A")}', styles['ChineseHeading1']))
        elements.append(Spacer(1, 0.1 * inch))
        image_path = os.path.join(images_folder, f'{question_id}.png')
        if os.path.exists(image_path):
            img = ImageReader(image_path)
            width, height = img.getSize()
            display_width = doc.width * 0.8
            elements.append(RLImage(image_path, width=display_width,
                                    height=display_width * height / float(width)))
        else:
            elements.append(Paragraph('（无图片）', normal_style))
        if export_answers:
            elements.append(Paragraph(f"答案：<font color='red'>{question.get('answer', 'N/A')}</font>", normal_style))
        elements.append(Spacer(1, 0.2 * inch))

    doc.build(elements)
